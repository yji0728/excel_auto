"""
Excel Operations Module
Provides core functionality for reading, writing, and manipulating Excel files.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import pandas as pd
from typing import List, Dict, Any, Optional


class ExcelOperations:
    """Class to handle Excel file operations"""
    
    def __init__(self):
        self.workbook = None
        self.worksheet = None
        self.file_path = None
    
    def open_file(self, file_path: str, sheet_name: Optional[str] = None) -> bool:
        """Open an Excel file"""
        try:
            self.file_path = file_path
            self.workbook = openpyxl.load_workbook(file_path)
            if sheet_name:
                self.worksheet = self.workbook[sheet_name]
            else:
                self.worksheet = self.workbook.active
            return True
        except Exception as e:
            print(f"Error opening file: {e}")
            return False
    
    def create_new_file(self) -> bool:
        """Create a new Excel workbook"""
        try:
            self.workbook = openpyxl.Workbook()
            self.worksheet = self.workbook.active
            return True
        except Exception as e:
            print(f"Error creating new file: {e}")
            return False
    
    def save_file(self, file_path: Optional[str] = None) -> bool:
        """Save the Excel file"""
        try:
            save_path = file_path if file_path else self.file_path
            if not save_path:
                print("No file path specified")
                return False
            self.workbook.save(save_path)
            return True
        except Exception as e:
            print(f"Error saving file: {e}")
            return False
    
    def read_cell(self, row: int, col: int) -> Any:
        """Read value from a specific cell"""
        try:
            return self.worksheet.cell(row=row, column=col).value
        except Exception as e:
            print(f"Error reading cell: {e}")
            return None
    
    def write_cell(self, row: int, col: int, value: Any) -> bool:
        """Write value to a specific cell"""
        try:
            self.worksheet.cell(row=row, column=col, value=value)
            return True
        except Exception as e:
            print(f"Error writing cell: {e}")
            return False
    
    def read_range(self, start_row: int, start_col: int, 
                   end_row: int, end_col: int) -> List[List[Any]]:
        """Read a range of cells"""
        try:
            data = []
            for row in range(start_row, end_row + 1):
                row_data = []
                for col in range(start_col, end_col + 1):
                    row_data.append(self.worksheet.cell(row=row, column=col).value)
                data.append(row_data)
            return data
        except Exception as e:
            print(f"Error reading range: {e}")
            return []
    
    def write_range(self, start_row: int, start_col: int, data: List[List[Any]]) -> bool:
        """Write data to a range of cells"""
        try:
            for i, row_data in enumerate(data):
                for j, value in enumerate(row_data):
                    self.worksheet.cell(row=start_row + i, column=start_col + j, value=value)
            return True
        except Exception as e:
            print(f"Error writing range: {e}")
            return False
    
    def apply_formatting(self, row: int, col: int, 
                        bold: bool = False, 
                        font_size: int = 11,
                        bg_color: Optional[str] = None,
                        align: str = 'left') -> bool:
        """Apply formatting to a cell"""
        try:
            cell = self.worksheet.cell(row=row, column=col)
            cell.font = Font(bold=bold, size=font_size)
            cell.alignment = Alignment(horizontal=align)
            if bg_color:
                cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            return True
        except Exception as e:
            print(f"Error applying formatting: {e}")
            return False
    
    def get_sheet_names(self) -> List[str]:
        """Get all sheet names in the workbook"""
        try:
            return self.workbook.sheetnames if self.workbook else []
        except Exception as e:
            print(f"Error getting sheet names: {e}")
            return []
    
    def add_sheet(self, sheet_name: str) -> bool:
        """Add a new sheet to the workbook"""
        try:
            self.workbook.create_sheet(title=sheet_name)
            return True
        except Exception as e:
            print(f"Error adding sheet: {e}")
            return False
    
    def delete_sheet(self, sheet_name: str) -> bool:
        """Delete a sheet from the workbook"""
        try:
            if sheet_name in self.workbook.sheetnames:
                del self.workbook[sheet_name]
                return True
            return False
        except Exception as e:
            print(f"Error deleting sheet: {e}")
            return False
    
    def to_dataframe(self) -> Optional[pd.DataFrame]:
        """Convert worksheet to pandas DataFrame"""
        try:
            data = self.worksheet.values
            cols = next(data)
            return pd.DataFrame(data, columns=cols)
        except Exception as e:
            print(f"Error converting to DataFrame: {e}")
            return None
    
    def from_dataframe(self, df: pd.DataFrame, start_row: int = 1, 
                      start_col: int = 1, include_header: bool = True) -> bool:
        """Write DataFrame to worksheet"""
        try:
            if include_header:
                for j, col_name in enumerate(df.columns):
                    self.worksheet.cell(row=start_row, column=start_col + j, value=col_name)
                start_row += 1
            
            for i, row in df.iterrows():
                for j, value in enumerate(row):
                    self.worksheet.cell(row=start_row + i, column=start_col + j, value=value)
            return True
        except Exception as e:
            print(f"Error writing DataFrame: {e}")
            return False
    
    def find_cells(self, search_value: Any, 
                   search_col: Optional[int] = None) -> List[tuple]:
        """Find cells containing a specific value"""
        try:
            results = []
            for row in self.worksheet.iter_rows():
                for cell in row:
                    if search_col and cell.column != search_col:
                        continue
                    if cell.value == search_value:
                        results.append((cell.row, cell.column))
            return results
        except Exception as e:
            print(f"Error finding cells: {e}")
            return []
    
    def close(self):
        """Close the workbook"""
        self.workbook = None
        self.worksheet = None
        self.file_path = None
